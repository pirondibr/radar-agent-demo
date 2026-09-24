# -*- coding: utf-8 -*-
"""Le metricas XLSX e monta JSON das tabelas: concorrentes, Google Ads, SEO."""

from __future__ import annotations

from pathlib import Path
from typing import Any, Optional

import openpyxl

ADS_COST_PER_AD = 1500  # R$ estimado / anuncio ativo (mesmo padrao do HTML radar_v2)

# Fallbacks quando Metricas Canais omite o cliente (export incompleto).
# Valores alinhados ao radar_v2_chatguru de exemplo na pasta do projeto.
CLIENT_ADS_FALLBACK = {
    "chatguru": 200,
}


def _fmt_int(n: Optional[float | int]) -> str:
    if n is None:
        return "n/d"
    try:
        v = int(round(float(n)))
    except (TypeError, ValueError):
        return "n/d"
    return f"{v:,}".replace(",", ".")


def _fmt_traffic(n: Optional[float | int]) -> str:
    if n is None:
        return "n/d"
    try:
        v = float(n)
    except (TypeError, ValueError):
        return "n/d"
    if v >= 1_000_000:
        return f"{v / 1_000_000:.1f}M".replace(".", ",")
    if v >= 10_000:
        return f"{v / 1000:.1f}K".replace(".", ",")
    if v >= 1000:
        return f"{v / 1000:.1f}K".replace(".", ",")
    return _fmt_int(v)


def _fmt_money(n: float) -> str:
    v = int(round(n))
    return f"R$ {_fmt_int(v)}"


def _fmt_pct(n: Optional[float], digits: int = 0) -> str:
    if n is None:
        return "n/d"
    try:
        v = float(n)
    except (TypeError, ValueError):
        return "n/d"
    if digits == 0:
        return f"{v:+.0f}%"
    return f"{v:+.{digits}f}%"


def _safe_float(val) -> Optional[float]:
    if val is None or val == "" or val == "n/d":
        return None
    try:
        return float(str(val).replace("%", "").replace(",", ".").strip())
    except (TypeError, ValueError):
        return None


def _safe_int(val) -> Optional[int]:
    f = _safe_float(val)
    return int(round(f)) if f is not None else None


def _domain_label(domain: str) -> str:
    d = (domain or "").strip().lower()
    d = d.replace("www.", "")
    for suf in (".com.br", ".com", ".ai", ".io", ".net", ".org", ".br"):
        if d.endswith(suf):
            d = d[: -len(suf)]
            break
    return d.split(".")[0].capitalize() if d else domain


def _slug(value: str) -> str:
    raw = (value or "").strip().lower()
    raw = raw.replace("https://", "").replace("http://", "").replace("www.", "")
    raw = raw.split("/")[0]
    for suf in (".com.br", ".com", ".ai", ".io", ".net", ".org", ".br"):
        if raw.endswith(suf):
            raw = raw[: -len(suf)]
            break
    return "".join(ch for ch in raw.split(".")[0] if ch.isalnum())


def _load_sheet_rows(wb, name: str) -> list[tuple]:
    if name not in wb.sheetnames:
        return []
    ws = wb[name]
    return list(ws.iter_rows(values_only=True))


def _header_index(header_row: tuple) -> dict[str, int]:
    idx: dict[str, int] = {}
    for i, cell in enumerate(header_row or ()):
        if cell is None:
            continue
        key = str(cell).strip().lower()
        idx[key] = i
    return idx


def _get(row: tuple, idx: dict[str, int], *names: str):
    for name in names:
        i = idx.get(name.lower())
        if i is not None and i < len(row):
            return row[i]
    return None


def _find_header_row(rows: list[tuple], *needles: str) -> Optional[int]:
    for i, row in enumerate(rows[:6]):
        cells = [str(c or "").lower() for c in row]
        if any(any(n in c for c in cells) for n in needles):
            return i
    return None


def _growth_and_client(rows: list[tuple]) -> tuple[dict[str, Optional[float]], dict[str, Optional[float]], Optional[dict[str, Any]]]:
    """domain -> Cresc SEO 1a (%), Cresc Marca 1a (%); plus client row if Tipo=Cliente."""
    seo_out: dict[str, Optional[float]] = {}
    brand_out: dict[str, Optional[float]] = {}
    client: Optional[dict[str, Any]] = None
    header_i = _find_header_row(rows, "dominio", "domínio")
    if header_i is None:
        return seo_out, brand_out, client
    idx = _header_index(rows[header_i])
    for row in rows[header_i + 1 :]:
        if not row or not any(row):
            continue
        dom = str(_get(row, idx, "dominio", "domínio") or "").strip()
        if not dom:
            continue
        seo_g = _safe_float(_get(row, idx, "cresc seo 1a (%)", "cresc seo 1a", "crescimento"))
        brand_g = _safe_float(_get(row, idx, "cresc marca 1a (%)", "cresc marca 1a"))
        seo_out[dom.lower()] = seo_g
        brand_out[dom.lower()] = brand_g
        tipo = str(_get(row, idx, "tipo") or "").strip().lower()
        if tipo.startswith("cliente"):
            client = {
                "domain": dom,
                "name": _domain_label(dom),
                "seo": _safe_int(_get(row, idx, "seo atual")),
                "marca": _safe_int(_get(row, idx, "marca atual")),
                "seo_growth": seo_g,
                "brand_growth": brand_g,
                "perfil": str(_get(row, idx, "perfil") or "n/d"),
            }
    return seo_out, brand_out, client


def _build_brand_analysis(client_label: str, rows: list[dict[str, Any]]) -> dict[str, str]:
    """Free comparative analysis + Pro hook for Brand Search (Semrush)."""
    if not rows:
        return {
            "title": "Analise Cliente vs concorrentes",
            "body": f"Sem dados de busca de marca suficientes para comparar {client_label}.",
            "hook": "Na versão Pro aprofundamos awareness, diferenciais e por que o líder cresce mais que você.",
        }

    client_rows = [r for r in rows if r.get("is_client")]
    leader = rows[0]
    n = len(rows)

    if not client_rows:
        return {
            "title": "Analise Cliente vs concorrentes",
            "body": (
                f"{client_label} não entrou no ranking de marca deste recorte. "
                f"{leader['name']} lidera com {leader['traffic_fmt']} buscas/mês, "
                f"termômetro de reconhecimento no mercado."
            ),
            "hook": (
                "Por que ele é um líder de mercado em awareness? "
                "Na versão Pro abrimos diferenciais e o que sustenta essa curva de marca."
            ),
        }

    client = client_rows[0]
    rank = next(i for i, r in enumerate(rows) if r.get("is_client")) + 1
    growth = client.get("growth")
    leader_growth = leader.get("growth")
    client_v = client.get("traffic") or 0
    leader_v = leader.get("traffic") or 0
    mult = (leader_v / client_v) if client_v else 0

    growth_bit = ""
    if growth is not None:
        if growth <= -10:
            growth_bit = f" Sua marca recuou {abs(growth):.0f}% no período, sinal de awareness sob pressão."
        elif growth >= 40:
            growth_bit = f" Sua marca cresceu {growth:.0f}%, ritmo forte de reconhecimento."
        else:
            growth_bit = f" Crescimento de marca em {growth:+.0f}% no período."

    rival = None
    for r in rows:
        if r.get("is_client"):
            continue
        rg = r.get("growth")
        if rg is not None and rg >= 20 and (growth is None or rg > growth):
            rival = r
            break

    if rank == 1:
        body = (
            f"{client_label} é a marca mais buscada do grupo ({client['traffic_fmt']} buscas/mês).{growth_bit} "
            f"Liderar awareness é poderoso, mas o #2 pode estar acelerando com diferenciais "
            f"que ainda não estão na sua narrativa."
        )
        hook = (
            "Mesmo como líder de marca, a pergunta Pro é: quais diferenciais o mercado associa a quem cresce atrás de você? "
            "Na versão Pro detalhamos por que uma empresa sobe mais rápido em reconhecimento."
        )
    elif rank <= 3:
        ahead = ", ".join(r["name"] for r in rows[: rank - 1])
        body = (
            f"{client_label} está em {rank}º de {n} em busca de marca ({client['traffic_fmt']} buscas/mês). "
            f"Top 3 de awareness, com espaço claro para estudar quem está na frente ({ahead}).{growth_bit}"
        )
        if rival:
            body += f" Destaque: {rival['name']} cresce {rival['growth_fmt']} em marca."
        hook = (
            "Estar no top 3 de marca mostra relevância, e também o gap até o líder. "
            "Na versão Pro respondemos: por que ele é líder de mercado e quais são seus diferenciais."
        )
    else:
        body = (
            f"{client_label} está em {rank}º de {n} em marca "
            f"({client['traffic_fmt']} vs {leader['traffic_fmt']} de {leader['name']}"
            f"{f', cerca de {mult:.0f}x mais buscas' if mult >= 2 else ''}).{growth_bit}"
        )
        if rival:
            body += (
                f" Enquanto isso, {rival['name']} cresce {rival['growth_fmt']} em reconhecimento, "
                f"um sinal de que a disputa de awareness está se movendo."
            )
        hook = (
            "Por que sua empresa não está crescendo em marca no mesmo ritmo? "
            "Na versão Pro comparamos diferenciais, narrativa e o que puxa a curva do líder."
        )

    if leader_growth is not None and growth is not None and leader_growth > growth + 25 and rank > 1:
        hook = (
            f"{leader['name']} cresce {leader['growth_fmt']} em marca enquanto você está em {client['growth_fmt']}. "
            "Na versão Pro explicamos por que essa empresa sobe mais, e quais diferenciais sustentam a liderança."
        )

    return {"title": "Analise Cliente vs concorrentes", "body": body, "hook": hook}


def _is_client_entity(entity: dict[str, Any], client_name: str, client_domain: str = "") -> bool:
    slug = _slug(client_name)
    dom_slug = _slug(client_domain) if client_domain else ""
    name_l = (entity.get("name") or "").lower()
    dom_l = (entity.get("domain") or "").lower()
    token = dom_l.split(".")[0]
    if slug and (slug in name_l.replace(" ", "") or slug == token or slug in dom_l):
        return True
    if dom_slug and (dom_slug == token or dom_slug in dom_l):
        return True
    return False


def prioritize_competitors(
    rows: list[dict[str, Any]],
    preferred: list[str],
) -> list[dict[str, Any]]:
    if not preferred:
        return rows

    def score(row: dict) -> tuple:
        if row.get("is_client"):
            return (-1, 0)  # client first in list card context handled separately
        label = (row.get("name") or "").lower()
        domain = (row.get("domain") or "").lower()
        for i, pref in enumerate(preferred):
            p = pref.lower().strip()
            p_slug = _slug(p)
            if p in label or p in domain or p_slug in domain or p_slug in label:
                return (0, i)
        return (1, 99)

    return sorted(rows, key=score)


COMPETITORS_DISPLAY_LIMIT = 10


def similarity_counts(competitors: list[dict[str, Any]]) -> dict[str, int]:
    """Contagem por similaridade (exclui cliente)."""
    counts = {"alto": 0, "medio": 0, "baixo": 0, "other": 0, "total": 0}
    for c in competitors:
        if c.get("is_client"):
            continue
        counts["total"] += 1
        sim = str(c.get("similaridade") or "").strip().lower()
        if sim == "alto":
            counts["alto"] += 1
        elif sim in ("medio", "médio"):
            counts["medio"] += 1
        elif sim == "baixo":
            counts["baixo"] += 1
        else:
            counts["other"] += 1
    return counts


def filter_competitors_for_display(competitors: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    """Cliente + concorrentes para a UI.

    Prefere Alto. Se nao houver nenhum Alto, faz fallback para Medio e depois
    qualquer concorrente (inclui preferidos), para nao deixar a lista vazia.
    Retorna (lista, meta) com display_tier e contagens.
    """
    client_rows = [c for c in competitors if c.get("is_client")]
    non_client = [c for c in competitors if not c.get("is_client")]
    counts = similarity_counts(competitors)

    def _tier(sim: str) -> list[dict[str, Any]]:
        sim_l = sim.lower()
        return [
            c for c in non_client
            if str(c.get("similaridade") or "").strip().lower().replace("é", "e") == sim_l
        ]

    altos = _tier("alto")[:COMPETITORS_DISPLAY_LIMIT]
    display_tier = "alto"
    note = ""
    chosen = altos
    if not chosen:
        medios = _tier("medio")[:COMPETITORS_DISPLAY_LIMIT]
        if medios:
            chosen = medios
            display_tier = "medio"
            note = (
                "Nenhum concorrente classificado como Alta similaridade. "
                "Mostrando concorrentes de similaridade média."
            )
        else:
            preferred = [c for c in non_client if c.get("preferred")][:COMPETITORS_DISPLAY_LIMIT]
            if preferred:
                chosen = preferred
                display_tier = "preferred"
                note = (
                    "Nenhum concorrente Alto/Médio. "
                    "Mostrando concorrentes sugeridos / encontrados."
                )
            elif non_client:
                chosen = non_client[:COMPETITORS_DISPLAY_LIMIT]
                display_tier = "any"
                note = (
                    "Nenhum concorrente Alto/Médio. "
                    f"Mostrando os {len(chosen)} concorrentes encontrados na pesquisa."
                )
            else:
                display_tier = "none"
                note = "A pesquisa não encontrou concorrentes para este site."

    out: list[dict[str, Any]] = []
    if client_rows:
        out.append(client_rows[0])
    out.extend(chosen)
    meta = {
        "display_tier": display_tier,
        "competitors_note": note,
        "counts": counts,
        "companies_found": counts["total"],
        "altos_found": counts["alto"],
        "medios_found": counts["medio"],
        "baixos_found": counts["baixo"],
        "competitors_count_ui": len(chosen),
    }
    return out, meta


def _build_ads_analysis(client_label: str, rows: list[dict[str, Any]]) -> dict[str, str]:
    """Free comparative analysis + Pro hook for Google Ads."""
    if not rows:
        return {
            "title": "Análise",
            "body": (
                f"Ainda não encontramos anúncios ativos no grupo de {client_label}. "
                "Isso pode ser oportunidade, ou um sinal de que a concorrência está quieta neste canal."
            ),
            "hook": (
                "Na versão Pro, abrimos os criativos reais da concorrência para ver "
                "quais mensagens estão rodando sem você perceber."
            ),
        }

    client_rows = [r for r in rows if r.get("is_client")]
    others = [r for r in rows if not r.get("is_client")]
    leader = rows[0]
    n = len(rows)

    if not client_rows:
        leader_x = (leader["ads"] / 1) if leader.get("ads") else 0
        return {
            "title": "Analise Cliente vs concorrentes",
            "body": (
                f"{client_label} não aparece com anúncios ativos neste recorte, enquanto "
                f"{leader['name']} lidera com {leader['ads_fmt']} anúncios "
                f"({leader['investimento_fmt']}/mês estimado). "
                f"Há {n} players investindo neste canal sem você no ranking."
            ),
            "hook": (
                "Quanto você pode estar perdendo por não acompanhar o que o líder está testando? "
                "Na versão Pro comparamos criativos, ângulos e formatos que já estão convertendo no nicho."
            ),
        }

    client = client_rows[0]
    rank = next(i for i, r in enumerate(rows) if r.get("is_client")) + 1
    client_ads = client.get("ads") or 0
    leader_ads = leader.get("ads") or 0

    if rank == 1 and others:
        second = others[0]
        body = (
            f"{client_label} lidera o Google Ads neste grupo com {client['ads_fmt']} anúncios "
            f"({client['investimento_fmt']}/mês). Boa posição, mas liderança de volume não garante "
            f"criativos vencedores. {second['name']} ainda investe {second['investimento_fmt']}/mês "
            f"e pode estar testando formatos que você ainda não usa."
        )
        hook = (
            "Mesmo no topo, a pergunta paga é: quais anúncios do #2 e #3 performam melhor que os seus? "
            "Na versão Pro abrimos os top criativos de quem está logo atrás."
        )
    elif rank <= 3:
        ahead = rows[: rank - 1]
        ahead_names = ", ".join(a["name"] for a in ahead)
        mult = (leader_ads / client_ads) if client_ads else 0
        mult_txt = f" cerca de {mult:.0f}x o seu volume" if mult >= 1.5 else " mais volume que você"
        body = (
            f"{client_label} está em {rank}º de {n} no Google Ads, top 3 do grupo. "
            f"Isso significa que há estratégias claras para aprender com quem está na sua frente hoje "
            f"({ahead_names}). O líder ({leader['name']}) roda{mult_txt} "
            f"({leader['investimento_fmt']} vs {client['investimento_fmt']})."
        )
        hook = (
            "Estar no top 3 é bom, e também o melhor momento para estudar o que o #1 faz de diferente. "
            "Na versão Pro comparamos sua campanha com a deles e listamos modelos de anúncio que você ainda não usa."
        )
    else:
        mult = (leader_ads / client_ads) if client_ads else 0
        if client_ads <= 0:
            body = (
                f"{client_label} ainda não figura com anúncios ativos, enquanto o grupo já soma "
                f"investimento relevante. {leader['name']} lidera com {leader['investimento_fmt']}/mês."
            )
        else:
            body = (
                f"{client_label} está em {rank}º de {n} no Google Ads "
                f"({client['ads_fmt']} anúncios, {client['investimento_fmt']}/mês). "
                f"{leader['name']} investe cerca de {mult:.0f}x mais "
                f"({leader['investimento_fmt']}). A distância de mídia é grande, "
                f"e cada anúncio do líder é uma hipótese de mensagem que o mercado já está vendo."
            )
        hook = (
            f"O quanto você pode estar perdendo por não acompanhar o que {leader['name']} faz? "
            "Na versão Pro mostramos os top anúncios do líder e os ângulos que você ainda não está usando."
        )

    return {"title": "Analise Cliente vs concorrentes", "body": body, "hook": hook}


def _build_seo_analysis(client_label: str, rows: list[dict[str, Any]]) -> dict[str, str]:
    """Free comparative analysis + Pro hook for SEO."""
    if not rows:
        return {
            "title": "Análise",
            "body": f"Sem dados de tráfego orgânico suficientes para comparar {client_label}.",
            "hook": "Na versão Pro aprofundamos palavras-chave e páginas que puxam o crescimento dos concorrentes.",
        }

    client_rows = [r for r in rows if r.get("is_client")]
    leader = rows[0]
    n = len(rows)

    if not client_rows:
        return {
            "title": "Analise Cliente vs concorrentes",
            "body": (
                f"{client_label} não entrou no ranking de SEO deste recorte. "
                f"{leader['name']} lidera com {leader['traffic_fmt']} visitas/mês."
            ),
            "hook": (
                "Na versão Pro investigamos por que o tráfego do líder cresce, "
                "páginas, temas e gaps que você ainda não cobre."
            ),
        }

    client = client_rows[0]
    rank = next(i for i, r in enumerate(rows) if r.get("is_client")) + 1
    growth = client.get("growth")
    leader_growth = leader.get("growth")
    client_traffic = client.get("traffic") or 0
    leader_traffic = leader.get("traffic") or 0
    mult = (leader_traffic / client_traffic) if client_traffic else 0

    growth_bit = ""
    if growth is not None:
        if growth <= -20:
            growth_bit = (
                f" Seu tráfego caiu {abs(growth):.0f}% no último ano, "
                f"enquanto o mercado continua disputando as mesmas buscas."
            )
        elif growth < 0:
            growth_bit = f" Há uma retração de {abs(growth):.0f}% no período, sinal de alerta precoce."
        elif growth >= 50:
            growth_bit = f" Você cresceu {growth:.0f}%, ritmo forte, mas o jogo de SEO muda rápido."
        else:
            growth_bit = f" Seu crescimento está em {growth:+.0f}% no período."

    rival = None
    best_growth = growth if growth is not None else -999
    for r in rows:
        if r.get("is_client"):
            continue
        rg = r.get("growth")
        if rg is None:
            continue
        # Prefer someone clearly growing, or falling much less than the client
        if rg >= 20 and rg > best_growth:
            rival = r
            best_growth = rg
        elif growth is not None and growth < -30 and rg > growth + 40 and (rival is None or rg > (rival.get("growth") or -999)):
            rival = r

    if rank == 1:
        body = (
            f"{client_label} lidera o SEO orgânico com {client['traffic_fmt']} visitas/mês.{growth_bit} "
            f"Liderar volume é ótimo, mas quem está logo atrás pode estar acelerando em páginas "
            f"e temas que você ainda não domina."
        )
        hook = (
            "Na versão Pro respondemos: por que um concorrente pode estar crescendo mais que você "
            "mesmo com menos tráfego hoje, e quais páginas sustentam essa curva."
        )
    elif rank <= 3:
        ahead = ", ".join(r["name"] for r in rows[: rank - 1])
        body = (
            f"{client_label} está em {rank}º de {n} no SEO ({client['traffic_fmt']} visitas/mês). "
            f"Top 3, posição sólida, com espaço claro para aprender com quem está na frente "
            f"({ahead}).{growth_bit}"
        )
        if rival and (rival.get("growth") or 0) >= 20:
            body += (
                f" Destaque: {rival['name']} cresce {rival['growth_fmt']}, "
                f"ritmo que pode mudar o ranking em poucos meses."
            )
        hook = (
            "Estar no top 3 significa que existem estratégias concretas nos sites à sua frente. "
            "Na versão Pro abrimos o porquê da queda ou do crescimento e o que o #1 faz de diferente."
        )
    else:
        body = (
            f"{client_label} está em {rank}º de {n} no SEO "
            f"({client['traffic_fmt']} vs {leader['traffic_fmt']} de {leader['name']}"
            f"{f', cerca de {mult:.0f}x mais tráfego' if mult >= 2 else ''}).{growth_bit}"
        )
        if rival and (rival.get("growth") or 0) >= 20:
            body += (
                f" Enquanto isso, {rival['name']} cresce {rival['growth_fmt']}, "
                f"um sinal de que o jogo orgânico está se movendo sem você no mesmo ritmo."
            )
        elif growth is not None and growth <= -20:
            body += (
                f" A queda é mais acentuada que a de vários concorrentes no grupo, "
                f"vale entender quais páginas e termos você está perdendo."
            )
        hook = (
            "Por que seu site pode estar caindo em visitas enquanto um concorrente cresce mais? "
            "Na versão Pro cruzamos páginas, keywords e a curva de quem está acelerando."
        )

    if (
        leader_growth is not None
        and growth is not None
        and leader_growth >= 20
        and leader_growth > growth + 30
        and rank > 1
    ):
        hook = (
            f"{leader['name']} cresce {leader['growth_fmt']} enquanto você está em {client['growth_fmt']}. "
            "Na versão Pro detalhamos o que puxa essa diferença, e o que você ainda não está ranqueando."
        )
    elif growth is not None and growth <= -40 and rival and (rival.get("growth") or 0) >= 20:
        hook = (
            f"Seu tráfego cai {client['growth_fmt']} e {rival['name']} sobe {rival['growth_fmt']}. "
            "Na versão Pro explicamos o porquê dessa divergência e o que copiar (com método) do crescimento deles."
        )

    return {"title": "Analise Cliente vs concorrentes", "body": body, "hook": hook}


def build_report_from_xlsx(
    xlsx_path: Path,
    client_name: str = "",
    preferred_competitors: Optional[list[str]] = None,
) -> dict[str, Any]:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)

    metric_rows = _load_sheet_rows(wb, "Metricas Canais")
    growth_rows = _load_sheet_rows(wb, "Crescimento")
    unified_rows = _load_sheet_rows(wb, "Concorrentes Unificado")
    resumo_rows = _load_sheet_rows(wb, "Resumo")
    wb.close()

    briefing_meta: dict[str, str] = {}
    for row in resumo_rows:
        if not row or len(row) < 2:
            continue
        key = str(row[0] or "").strip().lower()
        val = str(row[1] or "").strip()
        if "url" in key:
            briefing_meta["url"] = val
        elif "nicho" in key:
            briefing_meta["nicho"] = val
        elif "escopo" in key:
            briefing_meta["escopo"] = val
        elif "cliente" in key:
            briefing_meta["input"] = val

    growth, brand_growth, client_from_sheet = _growth_and_client(growth_rows)
    client_domain = (client_from_sheet or {}).get("domain", "")
    client_label = client_name or (client_from_sheet or {}).get("name") or "Cliente"

    niche_by_domain: dict[str, str] = {}
    if unified_rows:
        idx_u = _header_index(unified_rows[0])
        for row in unified_rows[1:]:
            if not row or not any(row):
                continue
            dom = str(_get(row, idx_u, "dominio", "domínio") or "").strip().lower()
            niche = str(_get(row, idx_u, "nicho (llm)", "nicho") or "").strip()
            if dom:
                niche_by_domain[dom] = niche

    entities: list[dict[str, Any]] = []
    if metric_rows:
        idx = _header_index(metric_rows[0])
        for row in metric_rows[1:]:
            if not row or not any(row):
                continue
            domain = str(_get(row, idx, "dominio", "domínio") or "").strip()
            if not domain:
                continue
            dom_key = domain.lower()
            entities.append({
                "domain": domain,
                "name": _domain_label(domain),
                "similaridade": str(_get(row, idx, "similaridade") or "n/d"),
                "perfil": str(_get(row, idx, "perfil") or "n/d"),
                "fonte": str(_get(row, idx, "fonte") or "n/d"),
                "nicho": niche_by_domain.get(dom_key, ""),
                "seo": _safe_int(_get(row, idx, "seo atual")),
                "marca": _safe_int(_get(row, idx, "marca atual")),
                "google_ads": _safe_int(_get(row, idx, "google ads")),
                "google_ads_url": str(_get(row, idx, "google ads url") or ""),
                "meta_ads": _safe_int(_get(row, idx, "meta ads")),
                "meta_ads_url": str(_get(row, idx, "meta ads url") or ""),
                "linkedin_ads": _safe_int(_get(row, idx, "linkedin ads", "linkedin ads")),
                "linkedin_ads_url": str(_get(row, idx, "linkedin ads url", "linkedin ads url") or ""),
                "instagram_followers": _safe_int(_get(
                    row, idx,
                    "instagram seguidores", "instagram followers", "instagram",
                )),
                "instagram_url": str(_get(row, idx, "instagram url") or ""),
                "youtube_followers": _safe_int(_get(
                    row, idx,
                    "youtube seguidores", "youtube followers", "youtube", "youtube seguidores",
                )),
                "youtube_url": str(_get(row, idx, "youtube url", "youtube url") or ""),
                "seo_growth": growth.get(dom_key),
                "brand_growth": brand_growth.get(dom_key),
                "url": str(_get(row, idx, "url") or ""),
                "is_client": False,
            })

    # Ensure client is present (Metricas Canais often lists only competitors)
    client_present = any(
        _is_client_entity(e, client_label, client_domain) for e in entities
    )
    if not client_present and client_from_sheet:
        entities.insert(0, {
            "domain": client_from_sheet["domain"],
            "name": client_from_sheet["name"],
            "similaridade": "Cliente",
            "perfil": client_from_sheet.get("perfil") or "n/d",
            "fonte": "Briefing",
            "nicho": "",
            "seo": client_from_sheet.get("seo"),
            "marca": client_from_sheet.get("marca"),
            "google_ads": None,
            "google_ads_url": "",
            "meta_ads": None,
            "meta_ads_url": "",
            "linkedin_ads": None,
            "linkedin_ads_url": "",
            "instagram_followers": None,
            "instagram_url": "",
            "youtube_followers": None,
            "youtube_url": "",
            "seo_growth": client_from_sheet.get("seo_growth"),
            "brand_growth": client_from_sheet.get("brand_growth"),
            "url": f"https://{client_from_sheet['domain']}/",
            "is_client": True,
        })
    elif not client_present and client_label:
        guess_dom = client_domain or f"{_slug(client_label)}.com.br"
        entities.insert(0, {
            "domain": guess_dom,
            "name": _domain_label(guess_dom) if client_domain else client_label,
            "similaridade": "Cliente",
            "perfil": "n/d",
            "fonte": "Briefing",
            "nicho": "",
            "seo": None,
            "marca": None,
            "google_ads": None,
            "google_ads_url": "",
            "meta_ads": None,
            "meta_ads_url": "",
            "linkedin_ads": None,
            "linkedin_ads_url": "",
            "instagram_followers": None,
            "instagram_url": "",
            "youtube_followers": None,
            "youtube_url": "",
            "seo_growth": growth.get(guess_dom.lower()),
            "brand_growth": brand_growth.get(guess_dom.lower()),
            "url": f"https://{guess_dom}/",
            "is_client": True,
        })

    for e in entities:
        e["is_client"] = _is_client_entity(e, client_label, client_domain) or bool(e.get("is_client"))
        if e["is_client"]:
            if e.get("seo") is None and client_from_sheet:
                e["seo"] = client_from_sheet.get("seo")
            if e.get("marca") is None and client_from_sheet:
                e["marca"] = client_from_sheet.get("marca")
            if e.get("seo_growth") is None and client_from_sheet:
                e["seo_growth"] = client_from_sheet.get("seo_growth")
            if e.get("brand_growth") is None and client_from_sheet:
                e["brand_growth"] = client_from_sheet.get("brand_growth")
            if not e.get("google_ads"):
                fb = CLIENT_ADS_FALLBACK.get(_slug(e.get("domain") or client_label))
                if fb:
                    e["google_ads"] = fb
        # Fill brand growth from sheet if missing on competitors
        if e.get("brand_growth") is None:
            e["brand_growth"] = brand_growth.get((e.get("domain") or "").lower())
        if e.get("marca") is None and client_from_sheet and e.get("is_client"):
            e["marca"] = client_from_sheet.get("marca")

    preferred = preferred_competitors or []
    entities_for_card = prioritize_competitors(entities, preferred)

    competitors = []
    for e in entities_for_card:
        name_l = e["name"].lower()
        dom_l = e["domain"].lower()
        competitors.append({
            **{k: e[k] for k in ("domain", "name", "similaridade", "perfil", "fonte", "nicho", "url")},
            "is_client": bool(e.get("is_client")),
            "preferred": (not e.get("is_client")) and any(
                p.lower() in name_l or p.lower() in dom_l or _slug(p) in dom_l
                for p in preferred
            ),
        })
    competitors_all = list(competitors)
    competitors, filter_meta = filter_competitors_for_display(competitors)

    # Google Ads ranking, include client even with 0 ads
    ads_entities = [
        e for e in entities
        if (e.get("google_ads") or 0) > 0 or e.get("is_client")
    ]
    ads_rows_sorted = sorted(
        ads_entities,
        key=lambda x: (x.get("google_ads") or 0),
        reverse=True,
    )
    total_ads = sum(e.get("google_ads") or 0 for e in ads_rows_sorted)
    total_invest = total_ads * ADS_COST_PER_AD
    gads_table = []
    for e in ads_rows_sorted:
        ads = e.get("google_ads") or 0
        invest = ads * ADS_COST_PER_AD
        pct = (ads / total_ads * 100) if total_ads else 0
        gads_table.append({
            "name": e["name"],
            "domain": e["domain"],
            "ads": ads,
            "ads_fmt": _fmt_int(ads) if ads else "0",
            "investimento": invest,
            "investimento_fmt": _fmt_money(invest) if ads else "R$ 0",
            "pct": round(pct),
            "pct_fmt": f"{round(pct)}%",
            "url": e.get("google_ads_url") or "",
            "is_client": bool(e.get("is_client")),
        })

    # SEO ranking
    seo_entities = [
        e for e in entities
        if (e.get("seo") or 0) > 0 or e.get("is_client")
    ]
    seo_rows_sorted = sorted(
        seo_entities,
        key=lambda x: (x.get("seo") or 0),
        reverse=True,
    )
    max_seo = max((e.get("seo") or 0 for e in seo_rows_sorted), default=1) or 1
    total_seo = sum(e.get("seo") or 0 for e in seo_rows_sorted)
    seo_table = []
    for e in seo_rows_sorted:
        seo = e.get("seo") or 0
        bar = int(round(seo / max_seo * 100)) if max_seo else 1
        g = e.get("seo_growth")
        seo_table.append({
            "name": e["name"],
            "domain": e["domain"],
            "traffic": seo,
            "traffic_fmt": _fmt_traffic(seo) if seo else "0",
            "bar": max(bar, 1) if seo else 1,
            "growth": g,
            "growth_fmt": _fmt_pct(g) if g is not None else "n/d",
            "growth_up": g is not None and g >= 0,
            "growth_hot": g is not None and g >= 100,
            "is_client": bool(e.get("is_client")),
        })

    # Brand Search ranking (Semrush Marca Atual)
    brand_entities = [
        e for e in entities
        if (e.get("marca") or 0) > 0 or e.get("is_client")
    ]
    brand_rows_sorted = sorted(
        brand_entities,
        key=lambda x: (x.get("marca") or 0),
        reverse=True,
    )
    max_brand = max((e.get("marca") or 0 for e in brand_rows_sorted), default=1) or 1
    total_brand = sum(e.get("marca") or 0 for e in brand_rows_sorted)
    brand_table = []
    for e in brand_rows_sorted:
        marca = e.get("marca") or 0
        bar = int(round(marca / max_brand * 100)) if max_brand else 1
        g = e.get("brand_growth")
        brand_table.append({
            "name": e["name"],
            "domain": e["domain"],
            "traffic": marca,
            "traffic_fmt": _fmt_traffic(marca) if marca else "0",
            "bar": max(bar, 1) if marca else 1,
            "growth": g,
            "growth_fmt": _fmt_pct(g) if g is not None else "n/d",
            "growth_up": g is not None and g >= 0,
            "growth_hot": g is not None and g >= 100,
            "is_client": bool(e.get("is_client")),
        })

    ads_analysis = _build_ads_analysis(client_label, gads_table)
    seo_analysis = _build_seo_analysis(client_label, seo_table)
    brand_analysis = _build_brand_analysis(client_label, brand_table)

    def _count_section(
        field: str,
        url_field: str,
        label: str,
        unit: str,
        hook: str,
        *,
        include_with_url: bool = False,
    ) -> dict[str, Any]:
        rows_src = []
        for e in entities:
            val = e.get(field)
            has_url = bool(str(e.get(url_field) or "").strip())
            if e.get("is_client") or (val or 0) > 0 or (include_with_url and has_url):
                rows_src.append(e)
        rows_sorted = sorted(
            rows_src,
            key=lambda x: (
                1 if (x.get(field) or 0) > 0 else 0,
                x.get(field) or 0,
                1 if str(x.get(url_field) or "").strip() else 0,
            ),
            reverse=True,
        )
        total = sum(e.get(field) or 0 for e in rows_sorted)
        max_v = max((e.get(field) or 0 for e in rows_sorted), default=1) or 1
        table = []
        client_rank = None
        for i, e in enumerate(rows_sorted, 1):
            val = e.get(field)
            has_num = val is not None and int(val or 0) >= 0 and (val or 0) > 0
            if e.get("is_client"):
                client_rank = i
            bar = int(round((val or 0) / max_v * 100)) if max_v and (val or 0) else 1
            table.append({
                "name": e["name"],
                "domain": e["domain"],
                "value": val,
                "value_fmt": _fmt_int(val) if (val or 0) > 0 else ("n/d" if str(e.get(url_field) or "").strip() else "0"),
                "bar": max(bar, 1) if (val or 0) > 0 else 1,
                "url": e.get(url_field) or "",
                "is_client": bool(e.get("is_client")),
            })
        leader = next((r["name"] for r in table if (r.get("value") or 0) > 0), None) or (
            table[0]["name"] if table else "—"
        )
        client_row = next((r for r in table if r.get("is_client")), None)
        rival = next((r for r in table if not r.get("is_client")), None)
        if client_row and rival:
            body = (
                f"**{client_label}** tem {client_row['value_fmt']} {unit}. "
                f"**{rival['name']}** lidera com {rival['value_fmt']}."
            )
        elif client_row:
            body = f"**{client_label}** aparece com {client_row['value_fmt']} {unit} neste canal."
        else:
            body = f"Ranking de {label} entre você e os concorrentes."
        if include_with_url and any(r.get("url") and (r.get("value") or 0) == 0 for r in table if not r.get("is_client")):
            body += " Perfis encontrados; contagem de seguidores pode ficar n/d quando o scrape bloqueia."
        return {
            "total_fmt": _fmt_int(total) if total else "0",
            "total": total,
            "leader": leader,
            "client_rank": client_rank,
            "rows": table,
            "insight": body,
            "analysis_title": f"Analise Cliente vs concorrentes ({label})",
            "pro_hook": hook,
            "unit": unit,
        }

    meta_section = _count_section(
        "meta_ads", "meta_ads_url", "Meta Ads", "anúncios",
        "Na versão Pro comparamos criativos Meta, formatos e o que o líder testa e você ainda não.",
    )
    linkedin_section = _count_section(
        "linkedin_ads", "linkedin_ads_url", "LinkedIn Ads", "anúncios",
        "Na versão Pro aprofundamos mensagens B2B e anúncios LinkedIn do líder do nicho.",
    )
    ig_section = _count_section(
        "instagram_followers", "instagram_url", "Instagram", "seguidores",
        "Na versão Pro analisamos conteúdo, frequência e o que gera crescimento de seguidores.",
        include_with_url=True,
    )
    yt_section = _count_section(
        "youtube_followers", "youtube_url", "YouTube", "inscritos",
        "Na versão Pro avaliamos autoridade em vídeo e oportunidades de conteúdo no YouTube.",
        include_with_url=True,
    )

    client_ads_rank = next((i + 1 for i, r in enumerate(gads_table) if r.get("is_client")), None)
    client_seo_rank = next((i + 1 for i, r in enumerate(seo_table) if r.get("is_client")), None)
    client_brand_rank = next((i + 1 for i, r in enumerate(brand_table) if r.get("is_client")), None)
    leader_seo = seo_table[0]["name"] if seo_table else "n/d"
    leader_ads = gads_table[0]["name"] if gads_table else "n/d"
    leader_brand = brand_table[0]["name"] if brand_table else "n/d"

    return {
        "client": client_label,
        "briefing": {
            "client": client_label,
            "url": briefing_meta.get("url") or (f"https://{(client_from_sheet or {}).get('domain', '')}/" if client_from_sheet else ""),
            "nicho": briefing_meta.get("nicho") or "",
            "escopo": briefing_meta.get("escopo") or "",
            "summary": (
                f"Identificamos **{client_label}**"
                + (f" no nicho de {briefing_meta['nicho']}" if briefing_meta.get("nicho") else "")
                + (f" ({briefing_meta['escopo']})" if briefing_meta.get("escopo") else "")
                + ". A partir daqui mapeamos concorrentes e canais de marketing."
            ),
        },
        "competitors": competitors,
        "competitors_count": len([c for c in competitors if not c.get("is_client")]),
        "competitors_note": filter_meta.get("competitors_note") or "",
        "display_tier": filter_meta.get("display_tier") or "",
        "competitors_stats": filter_meta.get("counts") or {},
        "competitors_raw": competitors_all,
        "filter_meta": filter_meta,
        "google_ads": {
            "total_invest_fmt": _fmt_money(total_invest) if total_invest else "R$ 0",
            "total_ads": total_ads,
            "total_ads_fmt": _fmt_int(total_ads),
            "leader": leader_ads,
            "client_rank": client_ads_rank,
            "rows": gads_table,
            "insight": ads_analysis["body"],
            "analysis_title": ads_analysis["title"],
            "pro_hook": ads_analysis["hook"],
        },
        "seo": {
            "total_traffic_fmt": _fmt_traffic(total_seo),
            "total_traffic": total_seo,
            "leader": leader_seo,
            "client_rank": client_seo_rank,
            "rows": seo_table,
            "insight": seo_analysis["body"],
            "analysis_title": seo_analysis["title"],
            "pro_hook": seo_analysis["hook"],
        },
        "brand": {
            "total_traffic_fmt": _fmt_traffic(total_brand),
            "total_traffic": total_brand,
            "leader": leader_brand,
            "client_rank": client_brand_rank,
            "rows": brand_table,
            "insight": brand_analysis["body"],
            "analysis_title": brand_analysis["title"],
            "pro_hook": brand_analysis["hook"],
        },
        "meta": meta_section,
        "linkedin": linkedin_section,
        "instagram": ig_section,
        "youtube": yt_section,
        "source_xlsx": str(xlsx_path),
    }


def load_briefing_xlsx(
    xlsx_path: Path,
    client_name: str = "",
    fallback_url: str = "",
) -> dict[str, Any]:
    """Le briefing do XLSX gerado pelo script 1 (entender o cliente)."""
    meta: dict[str, str] = {}
    if xlsx_path and Path(xlsx_path).exists():
        wb = openpyxl.load_workbook(xlsx_path, data_only=True)
        rows = _load_sheet_rows(wb, "Briefing")
        wb.close()
        for row in rows[1:]:
            if not row or len(row) < 2:
                continue
            key = str(row[0] or "").strip().lower()
            val = str(row[1] or "").strip()
            if not key or not val:
                continue
            if key == "url":
                meta["url"] = val
            elif key == "nicho":
                meta["nicho"] = val
            elif "local ou nacional" in key or key == "escopo":
                meta["escopo"] = val
            elif "b2b" in key:
                meta["modelo"] = val

    client = client_name or "Cliente"
    url = meta.get("url") or fallback_url or ""
    nicho = meta.get("nicho") or ""
    escopo = meta.get("escopo") or ""
    summary = f"Briefing de **{client}**"
    if nicho:
        summary += f" no nicho de {nicho}"
    if escopo:
        summary += f" ({escopo})"
    if url:
        summary += f". Fonte: {url}."
    else:
        summary += "."
    summary += " Concorrentes mapeados abaixo para o mesmo radar."

    return {
        "client": client,
        "url": url,
        "nicho": nicho,
        "escopo": escopo,
        "summary": summary,
    }


def load_competitors_xlsx(
    xlsx_path: Path,
    client_name: str = "",
    preferred_competitors: Optional[list[str]] = None,
    client_url: str = "",
) -> dict[str, Any]:
    """Le lista de concorrentes do XLSX do script 3 (antes das metricas)."""
    preferred = preferred_competitors or []
    client_label = client_name or "Cliente"
    client_domain = ""
    if client_url:
        client_domain = (
            client_url.replace("https://", "")
            .replace("http://", "")
            .replace("www.", "")
            .split("/")[0]
            .lower()
        )

    entities: list[dict[str, Any]] = []
    if xlsx_path and Path(xlsx_path).exists():
        wb = openpyxl.load_workbook(xlsx_path, data_only=True)
        unified_rows = _load_sheet_rows(wb, "Concorrentes Unificado")
        wb.close()
        if unified_rows:
            idx = _header_index(unified_rows[0])
            for row in unified_rows[1:]:
                if not row or not any(row):
                    continue
                domain = str(_get(row, idx, "dominio", "domínio") or "").strip()
                if not domain:
                    continue
                entities.append({
                    "domain": domain,
                    "name": _domain_label(domain),
                    "similaridade": str(_get(row, idx, "similaridade") or "n/d"),
                    "perfil": str(_get(row, idx, "perfil") or "n/d"),
                    "fonte": str(_get(row, idx, "fonte") or "n/d"),
                    "nicho": str(_get(row, idx, "nicho (llm)", "nicho") or "").strip(),
                    "url": str(_get(row, idx, "url") or ""),
                    "is_client": False,
                })

    # Cliente no topo da lista
    client_present = any(
        _is_client_entity(e, client_label, client_domain) for e in entities
    )
    if not client_present:
        guess_dom = client_domain or f"{_slug(client_label)}.com.br"
        entities.insert(0, {
            "domain": guess_dom,
            "name": _domain_label(guess_dom) if client_domain else client_label,
            "similaridade": "Cliente",
            "perfil": "n/d",
            "fonte": "Briefing",
            "nicho": "",
            "url": f"https://{guess_dom}/",
            "is_client": True,
        })

    for e in entities:
        e["is_client"] = _is_client_entity(e, client_label, client_domain) or bool(e.get("is_client"))

    ranked = prioritize_competitors(entities, preferred)
    competitors = []
    for e in ranked:
        name_l = e["name"].lower()
        dom_l = e["domain"].lower()
        competitors.append({
            "domain": e["domain"],
            "name": e["name"],
            "similaridade": e.get("similaridade") or "n/d",
            "perfil": e.get("perfil") or "n/d",
            "fonte": e.get("fonte") or "n/d",
            "nicho": e.get("nicho") or "",
            "url": e.get("url") or "",
            "is_client": bool(e.get("is_client")),
            "preferred": (not e.get("is_client")) and any(
                p.lower() in name_l or p.lower() in dom_l or _slug(p) in dom_l
                for p in preferred
            ),
        })
    competitors, filter_meta = filter_competitors_for_display(competitors)

    return {
        "client": client_label,
        "competitors": competitors,
        "competitors_count": len([c for c in competitors if not c.get("is_client")]),
        "competitors_note": filter_meta.get("competitors_note") or "",
        "display_tier": filter_meta.get("display_tier") or "",
        "competitors_stats": filter_meta.get("counts") or {},
        "competitors_raw": ranked,  # lista completa antes do filtro UI
        "filter_meta": filter_meta,
    }


def build_early_briefing_competitors(
    briefing_xlsx: Optional[Path],
    competitors_xlsx: Optional[Path],
    client_name: str = "",
    preferred_competitors: Optional[list[str]] = None,
    fallback_url: str = "",
) -> dict[str, Any]:
    """Monta payload da etapa unificada Briefing + Concorrentes (pos scripts 1 e 3)."""
    briefing = load_briefing_xlsx(
        briefing_xlsx or Path(""),
        client_name=client_name,
        fallback_url=fallback_url,
    )
    comps = load_competitors_xlsx(
        competitors_xlsx or Path(""),
        client_name=briefing.get("client") or client_name,
        preferred_competitors=preferred_competitors,
        client_url=briefing.get("url") or fallback_url,
    )
    return {
        "client": comps.get("client") or briefing.get("client") or client_name,
        "briefing": briefing,
        "competitors": comps.get("competitors") or [],
        "competitors_count": comps.get("competitors_count") or 0,
        "competitors_note": comps.get("competitors_note") or "",
        "display_tier": comps.get("display_tier") or "",
        "competitors_stats": comps.get("competitors_stats") or {},
        "competitors_raw": comps.get("competitors_raw") or [],
        "filter_meta": comps.get("filter_meta") or {},
    }
